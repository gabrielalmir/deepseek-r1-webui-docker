from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from structured_reporting.excel import ExcelAdapter
from structured_reporting.profile import ReportProfile


def make_profile(excel_value="None"):
    return ReportProfile.model_validate(
        {
            "version": 1,
            "name": "excel-demo",
            "base_prompt": "规则",
            "missing": {"excel_value": excel_value},
            "fields": [
                {"id": "category", "label": "类别", "type": "text", "rule": "规则"},
                {
                    "id": "length",
                    "label": "长度",
                    "excel_header": "长度(mm)",
                    "type": "number",
                    "rule": "规则",
                },
            ],
        }
    )


def create_workbook(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["case_id", "text", "existing"])
    sheet.append(["same", "first report", "keep-1"])
    sheet.append(["same", "second report", "keep-2"])
    sheet["A2"].font = Font(italic=True, color="123456")
    sheet["C2"].fill = PatternFill("solid", fgColor="DDEEFF")
    sheet.column_dimensions["A"].width = 19
    sheet.column_dimensions["C"].width = 31
    sheet.row_dimensions[2].height = 44
    workbook.save(path)


def test_reads_custom_sheet_and_header_names_with_source_row_numbers(tmp_path):
    source = tmp_path / "source.xlsx"
    create_workbook(source)

    rows = ExcelAdapter(
        sheet_name="Data", id_column="case_id", report_column="text"
    ).read(source)

    assert [(row.row_number, row.record_id, row.report) for row in rows] == [
        (2, "same", "first report"),
        (3, "same", "second report"),
    ]


def test_write_appends_after_last_header_and_preserves_source_format(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    create_workbook(source)
    before = load_workbook(source)["Data"]
    before_a2_style = copy(before["A2"]._style)
    before_c2_style = copy(before["C2"]._style)
    rows = ExcelAdapter(
        sheet_name="Data", id_column="case_id", report_column="text"
    ).read(source)

    ExcelAdapter(
        sheet_name="Data", id_column="case_id", report_column="text"
    ).write(
        source,
        output,
        profile=make_profile(),
        rows=rows,
        results=[
            {"类别": "A", "长度": 20.0},
            {"类别": None, "长度": None},
        ],
    )

    sheet = load_workbook(output)["Data"]
    assert [sheet.cell(1, column).value for column in range(1, 6)] == [
        "case_id",
        "text",
        "existing",
        "类别",
        "长度(mm)",
    ]
    assert [sheet.cell(2, column).value for column in range(4, 6)] == ["A", 20.0]
    assert [sheet.cell(3, column).value for column in range(4, 6)] == ["None", "None"]
    assert sheet["A2"]._style == before_a2_style
    assert sheet["C2"]._style == before_c2_style
    assert sheet.column_dimensions["A"].width == 19
    assert sheet.column_dimensions["C"].width == 31
    assert sheet.row_dimensions[2].height == 44
    assert sheet.freeze_panes == "D2"


def test_null_can_be_written_as_a_blank_cell(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    create_workbook(source)
    adapter = ExcelAdapter(sheet_name="Data", id_column="case_id", report_column="text")
    rows = adapter.read(source)

    adapter.write(
        source,
        output,
        profile=make_profile(excel_value=None),
        rows=rows[:1],
        results=[{"类别": None, "长度": None}],
    )

    sheet = load_workbook(output)["Data"]
    assert sheet["D2"].value is None
    assert sheet["E2"].value is None
