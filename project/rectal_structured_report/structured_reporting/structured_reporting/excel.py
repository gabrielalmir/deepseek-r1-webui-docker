from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Sequence

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .profile import ReportProfile


@dataclass(frozen=True)
class ExcelRecord:
    row_number: int
    record_id: str
    report: str


class ExcelAdapter:
    def __init__(
        self,
        *,
        sheet_name: str = "Sheet1",
        id_column: str = "patient_id",
        report_column: str = "Report",
    ) -> None:
        self.sheet_name = sheet_name
        self.id_column = id_column
        self.report_column = report_column

    def _load(self, path: str | Path):
        workbook = load_workbook(path)
        if self.sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(f"workbook must contain worksheet {self.sheet_name!r}")
        sheet = workbook[self.sheet_name]
        headers: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(1, column).value
            if value in (None, ""):
                continue
            name = str(value)
            if name in headers:
                workbook.close()
                raise ValueError(f"duplicate worksheet header: {name}")
            headers[name] = column
        missing = [
            name
            for name in (self.id_column, self.report_column)
            if name not in headers
        ]
        if missing:
            workbook.close()
            raise ValueError(f"worksheet is missing required headers: {', '.join(missing)}")
        return workbook, sheet, headers

    @staticmethod
    def _last_header_column(sheet) -> int:
        for column in range(sheet.max_column, 0, -1):
            if sheet.cell(1, column).value not in (None, ""):
                return column
        return 0

    @staticmethod
    def _last_source_row(sheet, id_column: int, report_column: int) -> int:
        for row_number in range(sheet.max_row, 1, -1):
            if sheet.cell(row_number, id_column).value not in (None, ""):
                return row_number
            if sheet.cell(row_number, report_column).value not in (None, ""):
                return row_number
        return 1

    def read(self, input_path: str | Path) -> list[ExcelRecord]:
        workbook, sheet, headers = self._load(input_path)
        id_index = headers[self.id_column]
        report_index = headers[self.report_column]
        last_row = self._last_source_row(sheet, id_index, report_index)
        rows = [
            ExcelRecord(
                row_number=row_number,
                record_id=""
                if sheet.cell(row_number, id_index).value is None
                else str(sheet.cell(row_number, id_index).value),
                report=""
                if sheet.cell(row_number, report_index).value is None
                else str(sheet.cell(row_number, report_index).value),
            )
            for row_number in range(2, last_row + 1)
        ]
        workbook.close()
        return rows

    @staticmethod
    def _excel_value(value: Any, null_value: str | None) -> Any:
        if value is None:
            return null_value
        if isinstance(value, (list, dict)):
            return json.dumps(value, ensure_ascii=False)
        return value

    def write(
        self,
        input_path: str | Path,
        output_path: str | Path,
        *,
        profile: ReportProfile,
        rows: Sequence[ExcelRecord],
        results: Sequence[dict[str, Any]],
    ) -> None:
        source = Path(input_path).resolve()
        destination = Path(output_path).resolve()
        if source == destination:
            raise ValueError("input and output paths must be different")
        if len(rows) != len(results):
            raise ValueError("rows and results must have the same length")

        workbook, sheet, _ = self._load(source)
        start_column = self._last_header_column(sheet) + 1
        output_labels = [field.output_label for field in profile.fields]
        excel_labels = [field.excel_label for field in profile.fields]

        header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        header_font = Font(bold=True, color="FFFFFF")
        header_alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        result_alignment = Alignment(vertical="top", wrap_text=True)

        for offset, label in enumerate(excel_labels):
            column = start_column + offset
            cell = sheet.cell(1, column)
            cell.value = label
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            sheet.column_dimensions[get_column_letter(column)].width = 22

        for row, result in zip(rows, results, strict=True):
            for offset, label in enumerate(output_labels):
                cell = sheet.cell(row.row_number, start_column + offset)
                cell.value = self._excel_value(
                    result.get(label), profile.missing.excel_value
                )
                cell.alignment = result_alignment

        sheet.freeze_panes = f"{get_column_letter(start_column)}2"
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_name(
            f".{destination.stem}.tmp{destination.suffix or '.xlsx'}"
        )
        try:
            workbook.save(temporary)
            workbook.close()
            temporary.replace(destination)
        finally:
            workbook.close()
            if temporary.exists():
                temporary.unlink()
